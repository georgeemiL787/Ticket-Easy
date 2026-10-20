"""Turn PDF, DOCX, XLSX and Markdown files into titled sections.

Each parser returns a list of Section(key, title, text). `key` becomes the citation anchor, so it must
be stable across re-ingestion: numbered headings use their number ("s2"), FAQ rows use their id.
"""

import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass
class Section:
    key: str
    title: str
    text: str


_NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.+)$")
_PDF_HEADING = re.compile(r"^(\d+)\.\s+[A-Z][^.]{0,70}$")


def _section_key(title: str, ordinal: int) -> str:
    match = _NUMBERED.match(title)
    return f"s{match.group(1)}" if match else f"s{ordinal}"


def _build(headed_blocks: list[tuple[str, list[str]]]) -> list[Section]:
    sections = []
    for ordinal, (title, lines) in enumerate(headed_blocks):
        body = "\n".join(line for line in lines if line.strip()).strip()
        if not body:
            continue
        sections.append(Section(_section_key(title, ordinal), title.strip(), body))
    return sections


_CITATION_LINE = re.compile(r"^\s*`[\w.-]+@[\w.-]+#[\w-]+`\s*$")  # a section's own citation, e.g. `refund_policy@v3#s1`
_HTML_TAG = re.compile(r"</?div\b[^>]*>")  # layout-only wrappers such as <div dir="rtl"> around Arabic text


def parse_markdown(path: Path) -> list[Section]:
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for line in path.read_text(encoding="utf-8").splitlines():
        if _CITATION_LINE.match(line):
            continue
        line = _HTML_TAG.sub("", line)
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading and len(heading.group(1)) >= 2:
            blocks.append((heading.group(2), []))
        elif heading:  # document title: keep it as the intro block's title
            blocks[0] = (heading.group(2), blocks[0][1])
        else:
            blocks[-1][1].append(line)
    # The intro block has no number; key it s0 so numbered sections keep their own numbers.
    return _build(blocks)


def parse_docx(path: Path) -> list[Section]:
    doc = Document(str(path))
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for para in doc.paragraphs:
        style = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading 1") or style == "title":
            blocks[0] = (text, blocks[0][1])
        elif style.startswith("heading"):
            blocks.append((text, []))
        else:
            blocks[-1][1].append(text)
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks[-1][1].extend(rows)
    return _build(blocks)


def parse_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    lines = []
    for page in reader.pages:
        lines.extend((page.extract_text() or "").splitlines())
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for line in lines:
        stripped = line.strip()
        if _PDF_HEADING.match(stripped):
            blocks.append((stripped, []))
        else:
            blocks[-1][1].append(stripped)
    # PDF text wraps mid-sentence; rejoin lines within a section.
    return [
        Section(s.key, s.title, re.sub(r"\s*\n\s*", " ", s.text))
        for s in _build(blocks)
    ]


def parse_xlsx(path: Path) -> list[Section]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    sections = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        id_col = headers.index("id") if "id" in headers else None
        for n, row in enumerate(rows[1:], start=2):
            values = {h: str(v).strip() for h, v in zip(headers, row) if v not in (None, "")}
            if not values:
                continue
            key = values.pop("id") if id_col is not None and "id" in values else f"{ws.title}-r{n}"
            text = "\n".join(f"{h}: {v}" for h, v in values.items())
            title = values.get("question_ar") or values.get("question_en") or key
            sections.append(Section(key.lower(), title, text))
    return sections


def parse_text(path: Path) -> list[Section]:
    """Plain text: a line like "2. Returns" starts a section, otherwise blank lines separate sections."""
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    blocks: list[tuple[str, list[str]]] = [("", [])]
    numbered = any(_NUMBERED.match(line) for line in raw.splitlines())
    for line in raw.splitlines():
        if numbered and _NUMBERED.match(line) and len(line) < 100:
            blocks.append((line.strip(), []))
        elif not numbered and not line.strip() and blocks[-1][1]:
            blocks.append(("", []))
        else:
            blocks[-1][1].append(line.strip())
    if not numbered:  # untitled paragraphs: the first line of each stands in for the title
        blocks = [((lines[0][:80] if lines else ""), lines) for _, lines in blocks]
    return _build(blocks)


class _HtmlSections(HTMLParser):
    """Headings (h1-h6) start sections; text of paragraphs, list items and table cells fills them. Scripts and styles are skipped."""

    SKIP = {"script", "style", "noscript", "template"}
    BLOCK = {"p", "li", "tr", "br", "div", "section", "article", "td", "th", "dd", "dt"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, list[str]]] = [("", [])]
        self.skip = 0
        self.heading: int | None = None
        self.buffer = ""

    def _flush(self) -> None:
        text = re.sub(r"\s+", " ", self.buffer).strip()
        self.buffer = ""
        if not text:
            return
        if self.heading is not None:
            if self.heading == 1:
                self.blocks[0] = (text, self.blocks[0][1])
            else:
                self.blocks.append((text, []))
        else:
            self.blocks[-1][1].append(text)

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = int(tag[1])
        elif tag in self.BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = None
        elif tag in self.BLOCK:
            self._flush()

    def handle_data(self, data):
        if not self.skip:
            self.buffer += data


def parse_html(path: Path) -> list[Section]:
    reader = _HtmlSections()
    reader.feed(path.read_text(encoding="utf-8-sig", errors="replace"))
    reader._flush()
    return _build(reader.blocks)


PARSERS = {
    ".txt": parse_text,
    ".html": parse_html,
    ".htm": parse_html,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".xlsx": parse_xlsx,
}


def parse(path: Path) -> list[Section]:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file type: {path.name} (supported: {', '.join(PARSERS)})")
    return parser(path)
