"""Generate the PDF, DOCX and XLSX seed documents for the fictional shop 'Nile Style' (shop_001).

The Markdown return policies are hand-written in data/corpus/shop_001. Everything here is fictional.
Run: python scripts/make_seed_docs.py
"""

from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "data" / "corpus" / "shop_001"

SHIPPING_EN = [
    ("1. Delivery areas",
     "Nile Style delivers to all 27 governorates of Egypt. We do not ship outside Egypt. "
     "Orders are shipped from our warehouse in Nasr City, Cairo."),
    ("2. Delivery times",
     "Estimated delivery time is 2-3 business days for Cairo and Giza, 3-5 business days for "
     "Alexandria, the Delta and Canal cities, and 5-7 business days for Upper Egypt, the Red Sea "
     "and Sinai. Business days are Saturday to Thursday, excluding official holidays."),
    ("3. Shipping fees",
     "Shipping costs EGP 50 for Cairo and Giza and EGP 70 for all other governorates. "
     "Shipping is free for orders of EGP 1,500 or more after discounts."),
    ("4. Cash on delivery",
     "Cash on delivery is available for orders up to EGP 5,000. Orders above EGP 5,000 must be "
     "paid in advance by card, mobile wallet or InstaPay."),
    ("5. Changing the delivery address",
     "The delivery address can be changed only while the order status is pending or processing. "
     "Once the order has been shipped the address cannot be changed; the customer may refuse the "
     "delivery and place a new order."),
    ("6. Late deliveries",
     "If an order arrives more than 3 business days after the latest estimated delivery date, the "
     "customer receives a EGP 100 voucher for the next order. Requests for compensation above this "
     "amount are reviewed by the customer care team and are not approved by the chat assistant."),
    ("7. Failed delivery attempts",
     "The courier attempts delivery twice and calls the customer before each attempt. After two "
     "failed attempts the order is returned to the warehouse and prepaid amounts are refunded "
     "minus the shipping fee."),
    ("8. Order cancellation",
     "An order can be cancelled free of charge before it is shipped. After shipping, cancellation "
     "is handled by the customer care team, and the customer may refuse the parcel at the door."),
    ("9. Order tracking",
     "A tracking link is sent by SMS and WhatsApp when the order is shipped. Customers can also "
     "track orders from the My Orders page using the order number, which starts with NS-."),
]

REFUND_AR = [
    ("1. متى يتم رد المبلغ",
     "يتم رد المبلغ بعد استلام المنتج المرتجع في المخزن وفحصه والتأكد من مطابقته لشروط الاسترجاع. "
     "لا يتم رد المبلغ لأي طلب مر على استلامه أكثر من 14 يومًا."),
    ("2. طرق رد المبلغ",
     "يتم رد المبلغ بنفس وسيلة الدفع الأصلية. البطاقات البنكية: خلال 7 إلى 14 يوم عمل حسب البنك. "
     "المحافظ الإلكترونية: خلال 3 أيام عمل. الطلبات المدفوعة عند الاستلام: يرد المبلغ على محفظة "
     "إلكترونية أو عبر انستاباي خلال 5 أيام عمل."),
    ("3. المبالغ الكبيرة",
     "طلبات رد المبالغ التي تزيد عن 3000 جنيه تحتاج إلى مراجعة وموافقة من فريق خدمة العملاء قبل "
     "التنفيذ، ولا يتم تنفيذها تلقائيًا."),
    ("4. ما لا يتم رده",
     "مصاريف الشحن الأصلية لا ترد إلا إذا كان المنتج تالفًا أو معيبًا أو تم إرسال منتج خاطئ. "
     "قيمة الكوبونات المستخدمة في الطلب لا ترد نقدًا، ويتم إصدار كوبون بنفس القيمة."),
    ("5. حالة الطلب المسموح فيها بالاسترداد",
     "يمكن رد المبلغ فقط للطلبات التي تم استلامها أو إرجاعها أو إلغاؤها قبل الشحن. الطلبات التي "
     "ما زالت في الطريق لا يتم رد مبلغها حتى تصل أو ترجع للمخزن."),
    ("6. المعاملات غير المعروفة والنزاعات",
     "إذا أبلغ العميل عن عملية دفع لم يقم بها أو اشتبه في استخدام بطاقته دون علمه، يتم تحويل "
     "الحالة فورًا إلى الفريق المختص ولا يتم التعامل معها عبر المساعد الآلي."),
]

FAQ = [
    ("q01", "ما هي مواعيد خدمة العملاء؟",
     "فريق خدمة العملاء متاح يوميًا من 10 صباحًا حتى 10 مساءً، والمساعد الآلي متاح 24 ساعة.",
     "What are your customer service hours?",
     "Our team is available daily from 10 am to 10 pm; the chat assistant is available 24/7."),
    ("q02", "ما هي طرق الدفع المتاحة؟",
     "نقبل البطاقات البنكية، والمحافظ الإلكترونية، وانستاباي، والدفع عند الاستلام للطلبات حتى 5000 جنيه.",
     "Which payment methods do you accept?",
     "Cards, mobile wallets, InstaPay, and cash on delivery for orders up to EGP 5,000."),
    ("q03", "كيف أختار المقاس المناسب؟",
     "يوجد جدول مقاسات في صفحة كل منتج. إذا كنت بين مقاسين ننصح باختيار المقاس الأكبر.",
     "How do I choose the right size?",
     "Each product page has a size chart. If you are between two sizes, choose the larger one."),
    ("q04", "هل لديكم فروع أو محل؟",
     "نايل ستايل متجر أونلاين فقط، ولا يوجد فرع لاستقبال العملاء.",
     "Do you have a physical store?",
     "Nile Style is online only; there is no walk-in store."),
    ("q05", "هل يوجد تغليف هدايا؟",
     "نعم، يمكن إضافة تغليف هدية مع كارت إهداء مقابل 30 جنيهًا للطلب.",
     "Do you offer gift wrapping?",
     "Yes, gift wrapping with a message card costs EGP 30 per order."),
    ("q06", "كيف أغير رقم الموبايل المسجل في حسابي؟",
     "لأسباب أمنية يتم تغيير رقم الموبايل بعد التحقق برمز يصل على الرقم القديم، أو عبر خدمة العملاء بعد التحقق من الهوية.",
     "How do I change the phone number on my account?",
     "For security, the phone number is changed after an OTP sent to the old number, or via customer care after identity checks."),
    ("q07", "هل يمكن استخدام أكثر من كوبون في نفس الطلب؟",
     "يمكن استخدام كوبون واحد فقط لكل طلب، ولا يمكن استخدام الكوبونات على منتجات التصفية.",
     "Can I use more than one voucher per order?",
     "Only one voucher per order, and vouchers cannot be used on clearance items."),
    ("q08", "هل تشحنون خارج مصر؟",
     "لا، نشحن داخل مصر فقط حاليًا.",
     "Do you ship internationally?",
     "No, we currently ship within Egypt only."),
    ("q09", "كيف أحصل على نقاط الولاء؟",
     "تحصل على نقطة واحدة لكل 10 جنيهات، وكل 100 نقطة تساوي كوبون خصم 25 جنيهًا. تنتهي النقاط بعد 12 شهرًا.",
     "How do loyalty points work?",
     "You earn 1 point per EGP 10; 100 points equal an EGP 25 voucher. Points expire after 12 months."),
    ("q10", "هل يمكن تعديل الطلب بعد تأكيده؟",
     "يمكن تعديل المنتجات أو المقاسات طالما حالة الطلب قيد المراجعة أو التجهيز، وبعد الشحن لا يمكن التعديل.",
     "Can I modify my order after confirming it?",
     "Items or sizes can be changed while the order is pending or processing; not after shipping."),
]


def make_pdf() -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, "Nile Style - Shipping and Delivery Policy", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, "Version 1, effective 1 September 2026.", new_x="LMARGIN", new_y="NEXT")
    for heading, body in SHIPPING_EN:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 8, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, body, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(OUT / "shipping_policy.pdf"))


def make_docx() -> None:
    doc = Document()
    doc.add_heading("سياسة رد المبالغ - نايل ستايل", level=1)
    doc.add_paragraph("النسخة الأولى، سارية من 1 سبتمبر 2026.")
    for heading, body in REFUND_AR:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
    doc.save(OUT / "refund_policy.docx")


def make_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "FAQ"
    ws.append(["id", "question_ar", "answer_ar", "question_en", "answer_en"])
    for row in FAQ:
        ws.append(list(row))
    wb.save(OUT / "faq.xlsx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf()
    make_docx()
    make_xlsx()
    print(f"Wrote shipping_policy.pdf, refund_policy.docx, faq.xlsx to {OUT}")
