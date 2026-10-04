"""Turn PDF, DOCX, XLSX and Markdown files into titled sections.

Each parser returns a list of Section(key, title, text). `key` becomes the citation anchor, so it must
be stable across re-ingestion: numbered headings use their number ("s2"), FAQ rows use their id.
"""

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass
class Section:
    key: str
    title: str
    text: str


_NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.+)$")
_PDF_HEADING = re.compile(r"^(\d+)\.\s+[A-Z][^.]{0,70}$")


def _section_key(title: str, ordinal: int) -> str:
    match = _NUMBERED.match(title)
    return f"s{match.group(1)}" if match else f"s{ordinal}"


def _build(headed_blocks: list[tuple[str, list[str]]]) -> list[Section]:
    sections = []
    for ordinal, (title, lines) in enumerate(headed_blocks):
        body = "\n".join(line for line in lines if line.strip()).strip()
        if not body:
            continue
        sections.append(Section(_section_key(title, ordinal), title.strip(), body))
    return sections


def parse_markdown(path: Path) -> list[Section]:
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for line in path.read_text(encoding="utf-8").splitlines():
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading and len(heading.group(1)) >= 2:
            blocks.append((heading.group(2), []))
        elif heading:  # document title: keep it as the intro block's title
            blocks[0] = (heading.group(2), blocks[0][1])
        else:
            blocks[-1][1].append(line)
    # The intro block has no number; key it s0 so numbered sections keep their own numbers.
    return _build(blocks)


def parse_docx(path: Path) -> list[Section]:
    doc = Document(str(path))
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for para in doc.paragraphs:
        style = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading 1") or style == "title":
            blocks[0] = (text, blocks[0][1])
        elif style.startswith("heading"):
            blocks.append((text, []))
        else:
            blocks[-1][1].append(text)
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks[-1][1].extend(rows)
    return _build(blocks)


def parse_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    lines = []
    for page in reader.pages:
        lines.extend((page.extract_text() or "").splitlines())
    blocks: list[tuple[str, list[str]]] = [("", [])]
    for line in lines:
        stripped = line.strip()
        if _PDF_HEADING.match(stripped):
            blocks.append((stripped, []))
        else:
            blocks[-1][1].append(stripped)
    # PDF text wraps mid-sentence; rejoin lines within a section.
    return [
        Section(s.key, s.title, re.sub(r"\s*\n\s*", " ", s.text))
        for s in _build(blocks)
    ]


def parse_xlsx(path: Path) -> list[Section]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    sections = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        id_col = headers.index("id") if "id" in headers else None
        for n, row in enumerate(rows[1:], start=2):
            values = {h: str(v).strip() for h, v in zip(headers, row) if v not in (None, "")}
            if not values:
                continue
            key = values.pop("id") if id_col is not None and "id" in values else f"{ws.title}-r{n}"
            text = "\n".join(f"{h}: {v}" for h, v in values.items())
            title = values.get("question_ar") or values.get("question_en") or key
            sections.append(Section(key.lower(), title, text))
    return sections


PARSERS = {
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".xlsx": parse_xlsx,
}


def parse(path: Path) -> list[Section]:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file type: {path.name} (supported: {', '.join(PARSERS)})")
    return parser(path)
